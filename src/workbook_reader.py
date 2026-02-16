from __future__ import annotations

from pathlib import Path


def read_formulas_from_workbook(path: Path) -> dict[str, str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required for workbook parsing. Install dependencies first.") from exc

    workbook = load_workbook(path, data_only=False)
    formulas: dict[str, str] = {}
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    formulas[f"{sheet.title}!{cell.coordinate}"] = value
    return formulas
